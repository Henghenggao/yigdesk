"""Generate the public Northwind FY2024 synthetic upload sample."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook


REVENUE_LINES = [
    ("Subscription", "Self-Serve MRR", [163.4, 793.8, 895.1, 613.8]),
    ("Subscription", "Sales-Led MRR", [292.7, 878.9, 754.1, 886.5]),
    ("Subscription", "Add-ons", [700.3, 255.4, 136.3, 101.1]),
    ("Usage", "API Overage", [470.3, 525.9, 592.7, 823.0]),
    ("Usage", "Storage Overage", [174.0, 757.0, 608.3, 716.3]),
    ("Services", "Onboarding", [391.3, 310.0, 775.9, 547.6]),
    ("Services", "Premium Support", [610.2, 94.2, 574.2, 215.9]),
]
COGS_LINES = [
    ("Infrastructure", "Compute", [-639.3, -332.0, -322.5, -697.4]),
    ("Infrastructure", "Storage", [-134.9, -524.1, -788.2, -442.9]),
    ("Infrastructure", "Bandwidth", [-465.9, -500.0, -785.9, -829.7]),
    ("Support & Ops", "Support Salaries", [-304.8, -490.0, -571.4, -514.8]),
    ("Support & Ops", "Payment Fees", [-759.4, -589.1, -92.4, -246.3]),
]


def generate(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    ledger = book.active
    ledger.title = "GL Detail"
    ledger["A1"] = "Northwind Cloud — Revenue & Cost Ledger Extract (SYNTHETIC)"
    ledger["A2"] = "Generated public demo data · no customer or production records"

    report = book.create_sheet("P&L Report")
    report.append(["Northwind Cloud — P&L FY2024 (USD k)"])
    report.append(["Category", "Account", "Line Item", "Q1", "Q2", "Q3", "Q4"])
    _append_category(report, "Revenue", REVENUE_LINES)
    _append_category(report, "COGS", COGS_LINES)

    notes = book.create_sheet("_Notes")
    notes["A1"] = "SYNTHETIC SAMPLE — safe for public demo and tests"
    notes["A2"] = "Quarterly P&L values are generated fixtures."
    book.save(path)
    book.close()
    return path


def _append_category(sheet, category: str, lines) -> None:
    sheet.append([category])
    current_account = None
    for account, item, quarters in lines:
        if account != current_account:
            sheet.append([None, account])
            current_account = account
        sheet.append([None, None, item, *quarters])


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("runtime/northwind-fy2024-synthetic.xlsx"),
    )
    args = parser.parse_args()
    output = generate(args.output.resolve())
    print(output)


if __name__ == "__main__":
    main()
