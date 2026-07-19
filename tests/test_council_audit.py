"""The blackboard ledger + committed DecisionRecord ARE the council audit.

This replaces the retired ``a2a`` post-hoc tool-sequence verifier. Instead of
re-checking a trace of read-only tool calls after the fact, we drive a real
decision through the deterministic blackboard on the ``council_discount``
scenario + policy and prove that:

  (a) the gate commits a ``DecisionRecord`` (chosen by ``max:headroom``, closed
      by policy, with grounded evidence + evaluator/source provenance);
  (b) the append-only ledger replays to the identical board + record, and the
      ``resolved`` op carries that record (the audit IS the op log); and
  (c) a missing required risk claim or missing CFO approval yields ``Pending``
      rather than a silent close.
"""

from __future__ import annotations

import json
from dataclasses import asdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from yigdesk.core.blackboard import Blackboard
from yigdesk.core.gate import Pending
from yigdesk.core.ledger import Ledger
from yigdesk.core.ops import RESOLVED
from yigdesk.core.projection import fold
from yigdesk.evaluator.expression import ExpressionEvaluator
from yigdesk.evaluator.model_source import ModelSource


ROOT = Path(__file__).resolve().parents[1]
SCENARIO = ROOT / "data" / "scenarios" / "council_discount"


def _model() -> dict:
    return json.loads((SCENARIO / "model.json").read_text(encoding="utf-8"))


def _policy() -> dict:
    return json.loads((SCENARIO / "policy.json").read_text(encoding="utf-8"))


def _headroom_after(candidate) -> str:
    """The priced 'after' headroom of a candidate's consequence."""

    return next(m.after for m in candidate.consequence.metrics if m.id == "headroom")


def _blackboard(tmp_path):
    """Blackboard on the real council_discount model + policy over a tmp ledger.

    The source workbook carries the canonical council figures (list ARR 1000k,
    COGS 480k, floor 40pt) so a 2% discount prices to +11.02pt headroom and a
    12% discount to +5.45pt — both feasible, so ``max:headroom`` has a real
    contest to decide.
    """

    model = _model()
    wb = Workbook()
    ws = wb.active
    ws.title = "Deal Inputs"
    ws["B2"] = 1000  # list_arr
    ws["B3"] = 0     # current discount
    ws["B4"] = 480   # cogs
    ws["B5"] = 40    # margin floor
    source_path = tmp_path / "council_deal.xlsx"
    wb.save(source_path)
    evaluator = ExpressionEvaluator(model)
    source = ModelSource(source_path, model["input_refs"])
    return Blackboard(tmp_path / "board.jsonl", evaluator, source), evaluator


def _drive(bb, *, with_claim: bool = True, with_approval: bool = True) -> None:
    """Open a council decision with two priced candidates, and (by default) the
    grounded risk claim and CFO approval the council policy requires.

    The higher-headroom candidate (c1, 2% discount) is proposed FIRST so a
    hypothetical "pick the last eligible candidate" selector bug can't quietly
    return the right answer.
    """

    bb.open_decision(
        "d1", "Approve the requested discount?", "council_discount", _policy(),
        actor="human:cfo", role="owner",
    )
    bb.propose_candidate(
        "d1", "c1", {"overrides": {"discount": 2}}, actor="agent:finance", role="proposer"
    )
    bb.propose_candidate(
        "d1", "c2", {"overrides": {"discount": 12}}, actor="agent:sales", role="proposer"
    )
    if with_claim:
        claim = bb.post_claim(
            "d1", "risk1", "risk", "d1",
            "A +5% COGS move erodes the thin discount headroom.",
            ["Deal Inputs!B4"],
            actor="agent:risk", role="critic",
        )
        # The claim must be grounded in a real evidence cell to count.
        assert claim.status == "grounded"
    if with_approval:
        bb.cast_approval("d1", "approve", "c1", actor="human:cfo", role="cfo")


def test_gate_commits_a_grounded_decision_record_chosen_by_policy(tmp_path):
    bb, evaluator = _blackboard(tmp_path)
    _drive(bb)

    candidates = bb.project().decisions["d1"].candidates
    # Both candidates price to "ok", so max:headroom is a genuine two-way
    # contest — a future floor/model change that makes one infeasible can't
    # silently degrade this into a single-candidate pick that still passes.
    assert candidates["c1"].consequence.verdict == "ok"
    assert candidates["c2"].consequence.verdict == "ok"
    # c1 (proposed first) is the strictly-higher-headroom option; the winner is
    # therefore neither the only eligible candidate nor the last-proposed one.
    assert Decimal(_headroom_after(candidates["c1"])) > Decimal(_headroom_after(candidates["c2"]))

    record = bb.request_resolve("d1", actor="human:cfo", role="cfo")

    assert not isinstance(record, Pending)
    # The policy closes on the genuinely-higher-headroom candidate.
    assert record.chosen_candidate_id == "c1"
    assert record.closed_by == "policy"
    assert record.rationale == "selector=max:headroom"
    # Grounded evidence + evaluator/source provenance are all in the record.
    chosen = candidates["c1"]
    assert record.evidence_refs == chosen.consequence.evidence_refs
    assert record.evidence_refs  # non-empty: the priced inputs are cited
    assert record.source_fingerprint == chosen.consequence.fingerprint
    assert record.evaluator_revision == evaluator.revision
    assert any(
        a["role"] == "cfo" and a["verdict"] == "approve" for a in record.approvals
    )
    assert bb.project().decisions["d1"].status == "resolved"


def test_append_only_ledger_replays_to_the_same_record(tmp_path):
    bb, _ = _blackboard(tmp_path)
    _drive(bb)

    record = bb.request_resolve("d1", actor="human:cfo", role="cfo")

    ops = Ledger(tmp_path / "board.jsonl").read()
    # The resolved op carries the full record: the ledger IS the audit trail.
    assert ops[-1].kind == RESOLVED
    assert ops[-1].payload["record"] == asdict(record)
    # Deterministic replay from disk reconstructs the identical board + record,
    # matching the live in-memory projection field for field.
    replay = fold(Ledger(tmp_path / "board.jsonl").read())
    decision = replay.decisions["d1"]
    assert decision.status == "resolved"
    assert decision.resolution == record
    assert replay == bb.project()


def test_missing_required_risk_claim_holds_instead_of_closing(tmp_path):
    bb, _ = _blackboard(tmp_path)
    # CFO approval present, but the policy-required grounded risk claim is absent.
    _drive(bb, with_claim=False)

    result = bb.request_resolve("d1", actor="human:cfo", role="cfo")

    assert isinstance(result, Pending)
    assert "risk" in result.reason
    # No silent close, and no resolved op is ever written.
    assert bb.project().decisions["d1"].status == "open"
    assert all(op.kind != RESOLVED for op in Ledger(tmp_path / "board.jsonl").read())


def test_missing_cfo_approval_holds_instead_of_closing(tmp_path):
    bb, _ = _blackboard(tmp_path)
    # Grounded risk claim present, but the required CFO approval is absent.
    _drive(bb, with_approval=False)

    result = bb.request_resolve("d1", actor="agent:risk", role="critic")

    assert isinstance(result, Pending)
    assert "approval" in result.reason
    assert bb.project().decisions["d1"].status == "open"
    assert all(op.kind != RESOLVED for op in Ledger(tmp_path / "board.jsonl").read())
