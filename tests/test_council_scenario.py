import json, subprocess, sys
from pathlib import Path
from yigdesk.evaluator.expression import ExpressionEvaluator
from yigdesk.evaluator.model_source import ModelSource
ROOT = Path(__file__).resolve().parents[1]

def _ev_src():
    subprocess.run([sys.executable, str(ROOT/"scripts/build_scenarios.py")], check=True)
    d = ROOT/"data/scenarios/council_discount"
    model = json.loads((d/"model.json").read_text())
    return ExpressionEvaluator(model), ModelSource(d/model["workbook"], model["input_refs"])

def test_council_discount_matches_engine_figures():
    ev, src = _ev_src()
    c = ev.price({"overrides":{"discount":12}}, src)
    m = {x.id: x for x in c.metrics}
    assert c.verdict == "ok"
    assert m["net_arr"].after == "880.00"
    assert m["gross_profit"].after == "400.00"
    assert m["gross_margin"].after == "45.45"
    assert m["headroom"].after == "5.45"

def test_council_discount_holds_without_cogs(tmp_path):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title="Deal Inputs"
    ws["B2"]=1000; ws["B3"]=0; ws["B5"]=40   # no B4/cogs
    p = tmp_path/"nocogs.xlsx"; wb.save(p)
    d = ROOT/"data/scenarios/council_discount"; model = json.loads((d/"model.json").read_text())
    c = ExpressionEvaluator(model).price({"overrides":{"discount":12}}, ModelSource(p, model["input_refs"]))
    assert c.verdict == "hold"
