from openpyxl import Workbook
from yigdesk.core.ledger import Ledger
from yigdesk.core.projection import fold
from yigdesk.core.blackboard import Blackboard
from yigdesk.core.gate import resolve
from yigdesk.evaluator.expression import ExpressionEvaluator
from yigdesk.evaluator.model_source import ModelSource

MODEL = {
    "input_refs": {"list_arr":"Deal Inputs!B2","discount":"Deal Inputs!B3","cogs":"Deal Inputs!B4","floor":"Deal Inputs!B5"},
    "metrics": [
        {"id":"net_arr","label":"Net ARR","formula":"list_arr * (1 - discount/100)","unit":"$k"},
        {"id":"gross_margin","label":"Gross margin","formula":"(net_arr - cogs) / net_arr * 100","requires":["cogs"],"unit":"%"},
        {"id":"headroom","label":"Headroom","formula":"gross_margin - floor","requires":["cogs"],"unit":"pt"},
    ],
    "constraints": [{"metric":"headroom","op":">=","value":0}],
}
def _bb(tmp_path):
    wb = Workbook(); ws = wb.active; ws.title="Deal Inputs"
    ws["B2"]=1000; ws["B3"]=10; ws["B4"]=480; ws["B5"]=40
    p = tmp_path/"deal.xlsx"; wb.save(p)
    return Blackboard(tmp_path/"board.jsonl", ExpressionEvaluator(MODEL), ModelSource(p, MODEL["input_refs"]))

def test_d2_replay_is_pure(tmp_path):
    bb = _bb(tmp_path)
    bb.open_decision("d1","q","discount",{"candidate_selector":"max:headroom"}, actor="h", role="owner")
    bb.propose_candidate("d1","c1",{"overrides":{"discount":12}}, actor="a", role="proposer")
    entries = Ledger(tmp_path/"board.jsonl").read()
    assert fold(entries) == fold(entries)
    assert repr(fold(entries)) == repr(fold(list(entries)))

def test_d1_same_action_same_consequence(tmp_path):
    bb = _bb(tmp_path)
    assert bb.ev.price({"overrides":{"discount":12}}, bb.src) == bb.ev.price({"overrides":{"discount":12}}, bb.src)

def test_d4_resolve_is_pure(tmp_path):
    bb = _bb(tmp_path)
    bb.open_decision("d1","q","discount",{"required_approvals":[{"role":"cfo","verdict":"approve"}],"candidate_selector":"max:headroom"}, actor="h", role="owner")
    bb.propose_candidate("d1","c1",{"overrides":{"discount":12}}, actor="a", role="proposer")
    bb.cast_approval("d1","approve","c1", actor="cfo", role="cfo")
    assert resolve(bb.project().decisions["d1"], "rev") == resolve(bb.project().decisions["d1"], "rev")

def test_d2_full_flow_replays_byte_identical_from_disk(tmp_path):
    import json as _json
    from dataclasses import asdict as _asdict
    bb = _bb(tmp_path)
    bb.open_decision("d1","q","discount",{"required_approvals":[{"role":"cfo","verdict":"approve"}],"candidate_selector":"max:headroom"}, actor="h", role="owner")
    bb.propose_candidate("d1","c1",{"overrides":{"discount":12}}, actor="a", role="proposer")
    bb.cast_approval("d1","approve","c1", actor="cfo", role="cfo")
    bb.request_resolve("d1", actor="cfo", role="cfo")
    def snapshot():
        board = fold(Ledger(tmp_path/"board.jsonl").read())
        d = board.decisions["d1"]
        return _json.dumps({"status": d.status, "record": _asdict(d.resolution)}, sort_keys=True)
    assert snapshot() == snapshot()
    assert bb.project().decisions["d1"].status == "resolved"
