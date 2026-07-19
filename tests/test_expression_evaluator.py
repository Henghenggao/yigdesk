from openpyxl import Workbook
from yigdesk.evaluator.model_source import ModelSource
from yigdesk.evaluator.expression import ExpressionEvaluator

MODEL = {
    "input_refs": {"list_arr":"Deal Inputs!B2","discount":"Deal Inputs!B3","cogs":"Deal Inputs!B4","floor":"Deal Inputs!B5"},
    "metrics": [
        {"id":"net_arr","label":"Net ARR","formula":"list_arr * (1 - discount/100)","unit":"$k"},
        {"id":"gross_margin","label":"Gross margin","formula":"(net_arr - cogs) / net_arr * 100","requires":["cogs"],"unit":"%"},
        {"id":"headroom","label":"Headroom","formula":"gross_margin - floor","requires":["cogs"],"unit":"pt"},
    ],
    "constraints": [{"metric":"headroom","op":">=","value":0}],
}

def _src(tmp_path, cogs=True):
    wb = Workbook(); ws = wb.active; ws.title = "Deal Inputs"
    ws["B2"]=1000; ws["B3"]=10; ws["B5"]=40
    if cogs: ws["B4"]=480
    p = tmp_path/"deal.xlsx"; wb.save(p)
    return ModelSource(p, MODEL["input_refs"])

def test_prices_candidate_with_before_after(tmp_path):
    ev = ExpressionEvaluator(MODEL)
    c = ev.price({"overrides":{"discount":12}}, _src(tmp_path))
    m = {x.id: x for x in c.metrics}
    assert c.verdict == "ok"
    assert m["net_arr"].before == "900.00" and m["net_arr"].after == "880.00"
    assert m["headroom"].after == "5.45"

def test_missing_required_input_yields_hold(tmp_path):
    ev = ExpressionEvaluator(MODEL)
    c = ev.price({"overrides":{"discount":12}}, _src(tmp_path, cogs=False))
    assert c.verdict == "hold"
    assert {x.id for x in c.metrics if x.after is None} >= {"gross_margin","headroom"}

def test_price_is_deterministic(tmp_path):
    ev = ExpressionEvaluator(MODEL); s = _src(tmp_path)
    assert ev.price({"overrides":{"discount":12}}, s) == ev.price({"overrides":{"discount":12}}, s)
