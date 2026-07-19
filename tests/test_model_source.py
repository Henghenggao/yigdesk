from openpyxl import Workbook
from yigdesk.evaluator.model_source import ModelSource

def _wb(tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = "Deal Inputs"
    ws["B2"] = 1000; ws["B3"] = 10
    p = tmp_path / "deal.xlsx"; wb.save(p); return p

def test_reads_named_inputs_and_reports_missing(tmp_path):
    src = ModelSource(_wb(tmp_path), {"list_arr":"Deal Inputs!B2","discount":"Deal Inputs!B3","cogs":"Deal Inputs!B4"})
    inputs = src.base_inputs()
    assert str(inputs["list_arr"]) == "1000" and str(inputs["discount"]) == "10"
    assert "cogs" not in inputs
    assert src.exists("Deal Inputs!B2") and not src.exists("Deal Inputs!B4")
    assert len(src.fingerprint) == 64
