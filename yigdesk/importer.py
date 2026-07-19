"""Strict synthetic XLSX ingestion for the public consequence-preview demo."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

import openpyxl
from openpyxl.utils import get_column_letter


PARSER_VERSION = "northwind-synthetic-pl/v1"
REQUIRED_SHEET = "P&L Report"
SYNTHETIC_MARKER = "SYNTHETIC"
MAX_UPLOAD_BYTES = 8 * 1024 * 1024
MAX_SHEETS = 32
MAX_ROWS = 100_000
MAX_COLUMNS = 128
MAX_ARCHIVE_MEMBERS = 2_048
MAX_UNCOMPRESSED_BYTES = 96 * 1024 * 1024


class WorkbookImportError(ValueError):
    """A safe, user-correctable workbook validation error."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class ImportedWorkbook:
    scenario: dict[str, Any]
    source: dict[str, Any]


def import_synthetic_workbook(
    path: Path,
    *,
    original_filename: str,
    requested_discount_pct: Decimal,
    margin_floor_pct: Decimal,
    current_discount_pct: Decimal = Decimal("0"),
) -> ImportedWorkbook:
    """Extract FY2024 Revenue and COGS from an uploaded synthetic P&L workbook."""

    _validate_percentage("requested_discount_pct", requested_discount_pct)
    _validate_percentage("current_discount_pct", current_discount_pct)
    _validate_percentage("margin_floor_pct", margin_floor_pct)
    if not original_filename.lower().endswith(".xlsx"):
        raise WorkbookImportError(
            "UNSUPPORTED_FILE_TYPE", "Choose a macro-free .xlsx workbook."
        )
    size = path.stat().st_size
    if size <= 0 or size > MAX_UPLOAD_BYTES:
        raise WorkbookImportError(
            "UPLOAD_SIZE_INVALID",
            f"Workbook must be between 1 byte and {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.",
        )
    _validate_xlsx_archive(path)
    before = _sha256(path)
    try:
        book = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as error:
        raise WorkbookImportError(
            "WORKBOOK_UNREADABLE", "The file is not a readable macro-free XLSX workbook."
        ) from error
    try:
        if len(book.sheetnames) > MAX_SHEETS:
            raise WorkbookImportError(
                "WORKBOOK_TOO_COMPLEX", f"Workbook may contain at most {MAX_SHEETS} sheets."
            )
        if REQUIRED_SHEET not in book.sheetnames:
            raise WorkbookImportError(
                "REQUIRED_SHEET_MISSING",
                f"The synthetic demo requires a '{REQUIRED_SHEET}' sheet.",
            )
        marker = str(book[book.sheetnames[0]]["A1"].value or "").upper()
        if SYNTHETIC_MARKER not in marker:
            raise WorkbookImportError(
                "SYNTHETIC_MARKER_MISSING",
                "This public demo accepts only generated workbooks marked SYNTHETIC in the first sheet.",
            )
        sheet_names = list(book.sheetnames)
        sheet = book[REQUIRED_SHEET]
        if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS:
            raise WorkbookImportError(
                "WORKBOOK_TOO_COMPLEX",
                f"The P&L sheet exceeds the {MAX_ROWS}-row or {MAX_COLUMNS}-column demo limit.",
            )
        extraction = _extract_pl(sheet)
    finally:
        book.close()
    after = _sha256(path)
    if before != after:
        raise WorkbookImportError(
            "SOURCE_DRIFT", "Workbook bytes changed while the read-only parser was running."
        )

    scenario = {
        "label": "Uploaded FY2024 evidence",
        "requester": "Elena Rossi",
        "requester_role": "Account Executive · synthetic",
        "recipient": "Marie Laurent · CFO",
        "subject": f"Northwind FY2024 — review {requested_discount_pct}% discount",
        "sent_at": "Demo session · just now",
        "message": (
            "Please test the proposed discount against FY2024 revenue, COGS, and the "
            "configured gross-margin floor before I reply."
        ),
        "list_arr_k": str(extraction["revenue_k"]),
        "current_discount_pct": str(current_discount_pct),
        "requested_discount_pct": str(requested_discount_pct),
        "cogs_k": str(extraction["cogs_k"]),
        "margin_floor_pct": str(margin_floor_pct),
    }
    source = {
        "kind": "uploaded-synthetic-xlsx",
        "filename": Path(original_filename).name,
        "size_bytes": size,
        "sha256": before,
        "parser_version": PARSER_VERSION,
        "synthetic_marker_verified": True,
        "analysis_bytes_unchanged": True,
        "sheet_names": sheet_names,
        "extraction": {
            "sheet": REQUIRED_SHEET,
            "period_columns": extraction["period_columns"],
            "revenue_cells": extraction["revenue_cells"],
            "cogs_cells": extraction["cogs_cells"],
            "revenue_k": str(extraction["revenue_k"]),
            "cogs_k": str(extraction["cogs_k"]),
        },
    }
    return ImportedWorkbook(scenario=scenario, source=source)


def parse_decimal_field(name: str, raw: Any) -> Decimal:
    try:
        value = Decimal(str(raw).strip())
    except (InvalidOperation, AttributeError, ValueError) as error:
        raise WorkbookImportError(
            "INVALID_DECISION_INPUT", f"{name} must be a decimal percentage."
        ) from error
    _validate_percentage(name, value)
    return value


def _extract_pl(sheet) -> dict[str, Any]:
    headers = [str(sheet.cell(2, column).value or "").strip() for column in range(1, 8)]
    if headers[:3] != ["Category", "Account", "Line Item"] or headers[3:] != [
        "Q1",
        "Q2",
        "Q3",
        "Q4",
    ]:
        raise WorkbookImportError(
            "PL_SCHEMA_MISMATCH",
            "P&L Report must use Category, Account, Line Item, Q1, Q2, Q3, Q4 headers.",
        )
    totals = {"Revenue": Decimal("0"), "COGS": Decimal("0")}
    addresses = {"Revenue": [], "COGS": []}
    detail_rows = {"Revenue": 0, "COGS": 0}
    category: str | None = None
    for row in range(3, sheet.max_row + 1):
        candidate = str(sheet.cell(row, 1).value or "").strip()
        if candidate:
            category = candidate
        if category not in totals or not sheet.cell(row, 3).value:
            continue
        detail_rows[category] += 1
        for column in range(4, 8):
            raw = sheet.cell(row, column).value
            try:
                value = Decimal(str(raw))
            except (InvalidOperation, ValueError) as error:
                raise WorkbookImportError(
                    "PL_VALUE_INVALID",
                    f"Expected a numeric quarterly value at {REQUIRED_SHEET}!{get_column_letter(column)}{row}.",
                ) from error
            totals[category] += value
            addresses[category].append(
                f"{REQUIRED_SHEET}!{get_column_letter(column)}{row}"
            )
    if detail_rows["Revenue"] == 0 or detail_rows["COGS"] == 0:
        raise WorkbookImportError(
            "PL_EVIDENCE_INCOMPLETE",
            "P&L Report must contain Revenue and COGS detail rows for FY2024.",
        )
    revenue = totals["Revenue"]
    cogs = abs(totals["COGS"])
    if revenue <= 0 or cogs <= 0:
        raise WorkbookImportError(
            "PL_EVIDENCE_INVALID", "Revenue and absolute COGS totals must be positive."
        )
    return {
        "period_columns": ["Q1", "Q2", "Q3", "Q4"],
        "revenue_k": revenue,
        "cogs_k": cogs,
        "revenue_cells": addresses["Revenue"],
        "cogs_cells": addresses["COGS"],
    }


def _validate_percentage(name: str, value: Decimal) -> None:
    if not value.is_finite() or value < 0 or value > 100:
        raise WorkbookImportError(
            "INVALID_DECISION_INPUT", f"{name} must be between 0 and 100."
        )


def _validate_xlsx_archive(path: Path) -> None:
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise WorkbookImportError(
                    "WORKBOOK_TOO_COMPLEX",
                    f"XLSX archive may contain at most {MAX_ARCHIVE_MEMBERS} members.",
                )
            uncompressed = sum(member.file_size for member in members)
            if uncompressed > MAX_UNCOMPRESSED_BYTES:
                raise WorkbookImportError(
                    "WORKBOOK_TOO_COMPLEX",
                    "The expanded XLSX archive exceeds the public demo safety limit.",
                )
            if any(
                member.filename.startswith(("/", "\\"))
                or ".." in Path(member.filename).parts
                for member in members
            ):
                raise WorkbookImportError(
                    "WORKBOOK_UNREADABLE", "The XLSX archive contains an unsafe member path."
                )
    except BadZipFile as error:
        raise WorkbookImportError(
            "WORKBOOK_UNREADABLE", "The file is not a valid XLSX archive."
        ) from error


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()
