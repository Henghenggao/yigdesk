"""Synthetic XLSX fixture creation and byte fingerprinting."""

from __future__ import annotations

import hashlib
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def fingerprint(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def create_workbook(path: Path, scenario: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    inputs = book.active
    inputs.title = "Deal Inputs"
    inputs.append(["Field", "Value", "Unit"])
    inputs.append(["List ARR", Decimal(scenario["list_arr_k"]), "$k"])
    inputs.append(["Discount", Decimal(scenario["current_discount_pct"]), "%"])
    inputs.append(["COGS", None if scenario["cogs_k"] is None else Decimal(scenario["cogs_k"]), "$k"])
    inputs.append(["Margin floor", Decimal(scenario["margin_floor_pct"]), "%"])
    inputs.append(["Requested discount", Decimal(scenario["requested_discount_pct"]), "%"])

    model = book.create_sheet("Deal Model")
    model.append(["Metric", "Value", "Formula definition"])
    model.append(["Net ARR", None, "List ARR × (1 − Discount)"])
    model.append(["Gross profit", None, "Net ARR − COGS"])
    model.append(["Gross margin", None, "Gross profit ÷ Net ARR"])
    model.append(["Margin headroom", None, "Gross margin − Margin floor"])

    for sheet in (inputs, model):
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="163D35")
            cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 20
        sheet.column_dimensions["C"].width = 34

    book.save(path)
    book.close()
